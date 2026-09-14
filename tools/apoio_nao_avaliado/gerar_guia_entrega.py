from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs" / "guia_entrega_dw_canada.docx"
NAVY = "18324B"
BLUE = "DCEAF7"
LIGHT = "F4F7FA"
GRAY = "D9E1E8"


def set_cell_shading(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_width(cell, inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_run(run, bold=False, size=11, color="000000", italic=False):
    run.bold = bold
    run.italic = italic
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)


def add_text(doc, text, bold=False, size=11, color="000000", align=None, before=0, after=7, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.22
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_run(run, bold=bold, size=size, color=color, italic=italic)
    return p


def h1(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_run(r, bold=True, size=16, color=NAVY)
    return p


def h2(doc, text):
    p = doc.add_paragraph(style="Heading 2")
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_run(r, bold=True, size=12, color=NAVY)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    set_run(r, size=10.5)


def numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    set_run(r, size=10.5)


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    hdr = t.rows[0]
    set_repeat_table_header(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        if widths:
            set_width(cell, widths[i])
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_run(r, bold=True, size=9.5, color="FFFFFF")
    for row_n, values in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(values):
            cell = cells[i]
            if widths:
                set_width(cell, widths[i])
            if row_n % 2:
                set_cell_shading(cell, LIGHT)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(value)
            set_run(r, size=9.3)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.8)
sec.bottom_margin = Inches(0.75)
sec.left_margin = Inches(0.85)
sec.right_margin = Inches(0.85)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
styles["Normal"].font.size = Pt(11)

# Cover
add_text(doc, "INSTITUTO FEDERAL DE ALAGOAS", bold=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER, before=60, after=3)
add_text(doc, "Campus Maceió · Sistemas de Informação", size=11, align=WD_ALIGN_PARAGRAPH.CENTER, after=85)
add_text(doc, "Guia de preparação", bold=True, size=19, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)
add_text(doc, "Data Warehouse de oportunidades para desenvolvedores imigrantes no Canadá", bold=True, size=22, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, after=18)
add_text(doc, "Conceitos, roteiro de execução, revisão da Entrega 1 e proposta de apresentação", size=12, align=WD_ALIGN_PARAGRAPH.CENTER, after=96)
add_text(doc, "Equipe: Eliezir Moreira Peixoto Neto · Raphael Phillipe da Silva Silverio", size=11, align=WD_ALIGN_PARAGRAPH.CENTER, after=5)
add_text(doc, "Maceió - AL · 14 de dezembro de 2026", size=11, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
doc.add_page_break()

h1(doc, "1. Como usar este guia")
add_text(doc, "Este documento organiza o que a equipe precisa saber e fazer para apresentar a primeira entrega do projeto. Ele não substitui o trabalho escrito: serve como roteiro de estudo, preparação da fala e checklist de correções antes do envio.")
h2(doc, "Resultado esperado da atividade")
add_text(doc, "A equipe deve demonstrar que consegue transformar dados operacionais de fontes diferentes em um modelo dimensional adequado para análise. O foco da primeira entrega é o desenho do Data Warehouse, a documentação das fontes e a justificativa técnica das transformações - não a implementação completa de dashboards ou ETL em produção.")
h2(doc, "O projeto em uma frase")
add_text(doc, "O Data Warehouse reúne indicadores oficiais de salário, imigração brasileira e habitação para comparar as províncias e territórios canadenses sob a perspectiva de uma pessoa desenvolvedora de software que deseja imigrar.", bold=True, size=11)

h1(doc, "2. Conceitos essenciais")
h2(doc, "O que é um Data Warehouse")
add_text(doc, "Um Data Warehouse (DW) é um repositório de dados organizado para consulta, comparação histórica e apoio à decisão. Em vez de registrar transações do dia a dia, ele consolida dados de fontes diferentes em uma estrutura estável e compreensível para análise.")
table(doc, ["Conceito", "Explicação aplicada ao projeto"], [
    ("Dados de origem / OLTP", "São os CSVs e tabelas públicas de onde os dados vêm. Neste projeto, eles representam a camada operacional de origem, embora não sejam sistemas transacionais internos."),
    ("ETL", "Extract, Transform, Load: extrair arquivos oficiais, aplicar regras documentadas e carregar dimensões e fatos no Oracle."),
    ("Staging", "Área intermediária que espelha os arquivos de origem. É onde filtros, validações e padronizações são feitos sem destruir os arquivos brutos."),
    ("Dimensão", "Tabela que descreve o contexto da análise: província, ocupação, período, país de nascimento, tipo de moradia e fonte."),
    ("Tabela fato", "Tabela com as medidas que serão analisadas: salários, quantidade de imigrantes e indicadores de habitação."),
    ("Star Schema", "Modelo em que uma tabela fato se conecta às dimensões. Ele simplifica consultas e evita misturar medidas de granularidades diferentes."),
    ("Surrogate key", "Chave numérica interna do DW. Ela preserva as chaves originais como referência e torna as relações mais controladas."),
], [1.55, 4.85])

h2(doc, "Por que existem três fatos")
add_text(doc, "Salário, imigração e habitação têm medidas e períodos diferentes. Colocar tudo em uma única tabela fato repetiria salários em cada período de imigração e poderia gerar somas incorretas. Por isso o projeto usa fato_salario, fato_imigracao e fato_habitacao, ligadas por dimensões conformadas, principalmente dim_provincia.")
h2(doc, "Cuidado com o tempo")
add_text(doc, "Imigração e habitação são dados do Censo de 2021; salários têm período de referência 2023-2024. O DW permite compará-los como indicadores, mas a apresentação não deve dizer que todos pertencem ao mesmo ano nem afirmar causalidade.")

h1(doc, "3. Dados e modelo deste projeto")
table(doc, ["Fonte", "O que fornece", "Uso analítico"], [
    ("Job Bank / ESDC", "Salário mínimo, mediano, máximo, médio e quartis para Software developers and programmers, NOC 21232.", "Comparar remuneração por província ou região disponível."),
    ("Statistics Canada 98-10-0307-01", "Imigrantes por local de nascimento e período de imigração no Censo 2021.", "Medir comunidade brasileira e sua distribuição temporal."),
    ("Statistics Canada 98-10-0258-01", "Indicadores de moradia por condição de posse no Censo 2021.", "Avaliar a proporção de inquilinos com custo de moradia inacessível."),
], [1.75, 2.75, 1.9])
h2(doc, "Fluxo dos dados")
table(doc, ["Etapa", "O que acontece", "Regra de qualidade"], [
    ("1. Landing", "Guardar o download original, URL, data de extração e checksum.", "Nunca alterar ou apagar o arquivo bruto."),
    ("2. Staging", "Criar tabelas que refletem os CSVs como texto e validar campos.", "Preservar valores ausentes como N/A ou NULL."),
    ("3. Transformação", "Filtrar NOC 21232, Brazil, totais de idade/gênero, 13 jurisdições e Censo 2021; normalizar nomes e códigos.", "Registrar a justificativa e o impacto de cada filtro."),
    ("4. DW", "Carregar dimensões, chaves substitutas e as três tabelas fato.", "Aplicar PK, FK, AK, CHECK, NOT NULL, índices e sequências."),
    ("5. BI", "Executar consultas e gráficos para responder às 10 perguntas.", "Distinguir dado indisponível de valor zero ou estimado."),
], [1.1, 3.25, 2.05])

h1(doc, "4. Passo a passo para realizar a atividade")
numbered(doc, "Confirmar os requisitos da entrega, guardar a autorização para trabalho em dupla e definir os papéis da equipe.")
numbered(doc, "Delimitar o problema: escolher províncias e territórios canadenses para comparar oportunidades de imigração para desenvolvedores de software.")
numbered(doc, "Documentar as fontes, períodos e limitações. Explicar que os CSVs públicos são bases operacionais de origem, não OLTP transacional interno.")
numbered(doc, "Fazer o dicionário de dados da origem: listar campos, tipos, domínio e finalidade de cada arquivo usado.")
numbered(doc, "Executar a análise exploratória: quantidade de registros, valores ausentes, duplicidades, faixas salariais, distribuição da comunidade brasileira e indicador habitacional.")
numbered(doc, "Documentar cada mudança: filtro aplicado, motivo, quantidade antes/depois e impacto na análise. Nenhum bruto deve ser sobrescrito.")
numbered(doc, "Definir o grão de cada fato antes de desenhar as tabelas. Exemplo: fato_salario = província + ocupação + período salarial.")
numbered(doc, "Desenhar e explicar o Star Schema, dicionário do DW, constraints, sequências e tablespaces.")
numbered(doc, "Escrever 10 perguntas respondíveis e demonstrar qual fato, dimensão e medida sustentam cada uma.")
numbered(doc, "Preparar a apresentação, testar o script Oracle e revisar os arquivos com nomes e referências consistentes.")

h1(doc, "5. Ferramentas utilizadas e recomendadas")
table(doc, ["Ferramenta", "Função", "Como explicar em sala"], [
    ("Oracle Database / SQL Developer", "Criar o DW, tablespaces, sequências, tabelas, constraints e consultas SQL.", "É o destino dimensional e garante integridade dos dados."),
    ("Python", "Ler CSVs, executar EDA e gerar gráficos para o relatório e a apresentação.", "Automatiza a análise reprodutível sem alterar os dados brutos."),
    ("Pentaho Data Integration", "Implementar futuramente as transformações e cargas ETL idempotentes.", "É a ferramenta visual para levar dados de staging ao DW."),
    ("Oracle SQL Data Modeler", "Desenhar o modelo lógico/físico e o Star Schema.", "Ajuda a visualizar fatos, dimensões e relações."),
    ("HTML, CSS e JavaScript", "Executar o deck interativo já existente no repositório.", "Permite apresentar como slides e demonstrar gráficos no navegador."),
], [1.65, 2.55, 2.2])

h1(doc, "6. Revisão da entrega atual")
add_text(doc, "A avaliação abaixo considera os requisitos fornecidos para a Entrega 1 e o conteúdo existente no repositório.")
table(doc, ["Item exigido", "Situação", "Evidência ou correção necessária"], [
    ("Trabalho em trio", "DUPLA COM AUTORIZAÇÃO A CONFIRMAR", "O projeto será apresentado por Eliezir Moreira Peixoto Neto e Raphael Phillipe da Silva Silverio. As instruções indicam trio; guardar a autorização formal do professor para a dupla."),
    ("Duas ou mais bases", "CONFORME", "Há duas fontes institucionais: Statistics Canada e Job Bank/ESDC, com três conjuntos de dados."),
    ("Objetivo e meta SMART", "CONFORME", "O objetivo está documentado e a meta SMART possui prazo definido: 14 de dezembro de 2026."),
    ("Dez perguntas", "CONFORME", "As 10 perguntas estão listadas no plano, relatório e apresentação."),
    ("Slides", "PARCIAL", "Existe deck HTML com 13 telas. Recomenda-se exportar também uma versão PDF para evitar problema de compatibilidade."),
    ("Trabalho escrito", "CONFORME", "docs/dw_canada.docx cobre escopo, fontes, EDA, alterações, Star Schema, dicionário do DW e anexo."),
    ("Script Oracle", "PARCIAL", "O DDL foi revisado estaticamente: ordem, constraints e identificadores estão consistentes. Falta executar no Oracle do ambiente da disciplina, incluindo a seção DBA de tablespaces."),
    ("Consistência do repositório", "CONFORME", "Referências de arquivos e a grafia do nome Raphael Phillipe foram padronizadas nos materiais do projeto."),
    ("Qualidade visual do relatório", "CONFORME", "A quebra de página desnecessária foi removida e o documento foi revisado visualmente."),
], [1.45, 1.05, 3.9])

h2(doc, "Prioridade de correção")
bullet(doc, "Crítica: confirmar e guardar a autorização para entrega em dupla. Sem essa autorização, a composição diverge da instrução explícita de trabalho em trio.")
bullet(doc, "Alta: testar o script SQL no Oracle da disciplina, se o ambiente estiver disponível, e registrar qualquer ajuste específico de tablespace." )
bullet(doc, "Média: gerar PDF dos slides como cópia de segurança para a apresentação.")
bullet(doc, "Planejamento: o README indica que ETL, dashboards e grafos estão pendentes. Isso não invalida a Entrega 1, mas não deve ser apresentado como projeto completo finalizado.")

h1(doc, "7. Forma recomendada de apresentar em aula")
add_text(doc, "Recomendação: usar o deck HTML existente como uma apresentação interativa de slides, abrindo presentation/index.html em tela cheia. Assim a equipe mantém a estrutura de slides exigida e pode mostrar gráficos e dados reais no navegador. Levar também um PDF exportado como plano B.")
doc.add_page_break()
h2(doc, "Roteiro de 10 a 12 minutos")
table(doc, ["Tempo", "Tela / assunto", "Mensagem principal"], [
    ("0:00-0:45", "Abertura", "Apresentar tema, equipe e pergunta central: onde um desenvolvedor imigrante encontra melhor equilíbrio entre salário, comunidade e moradia?"),
    ("0:45-2:00", "Problema e objetivo", "Os indicadores estavam em fontes diferentes; o DW os organiza para análise comparável."),
    ("2:00-3:00", "Fontes e limites", "São duas fontes institucionais e três conjuntos de dados; Censo é 2021 e salário é 2023-2024. Não há causalidade nem estimativa de ausência."),
    ("3:00-4:30", "As 10 perguntas", "Mostrar que as perguntas são o requisito do negócio e orientam o modelo."),
    ("4:30-6:30", "EDA", "Mostrar um gráfico de salário, um de comunidade brasileira e um de habitação. Explicar o que cada um revela."),
    ("6:30-8:30", "Star Schema", "Explicar dimensões, fatos, grão e por que existem três fatos."),
    ("8:30-9:45", "Tratamento dos dados", "Explicar filtros, preservação de bruto, ausência como NULL/N/A e impacto dos recortes."),
    ("9:45-11:00", "Entrega e próximos passos", "Apontar o SQL Oracle, dicionários e próxima etapa: ETL e dashboards."),
], [0.8, 1.7, 3.9])
h2(doc, "Divisão de fala para a dupla")
bullet(doc, "Pessoa 1: contexto, problema, objetivo, meta SMART, fontes e análise exploratória.")
bullet(doc, "Pessoa 2: qualidade dos dados, filtros, Star Schema, dicionário DW, script Oracle, perguntas e próximos passos.")
h2(doc, "Demonstração que vale a pena")
add_text(doc, "Durante a fala, abra somente três gráficos já existentes: salário mediano, comunidade brasileira e habitação de inquilinos. Em seguida mostre o diagrama do modelo. A demonstração deve explicar uma decisão, não apenas navegar pelas telas.")

h1(doc, "8. Checklist antes da aula")
table(doc, ["Verificação", "Pronto?"], [
    ("Autorização formal do professor para trabalho em dupla guardada", "[  ]"),
    ("Nome dos integrantes igual em todos os arquivos", "[  ]"),
    ("Data exata inserida na meta SMART", "[  ]"),
    ("Script Oracle testado sem linhas de separador inválidas", "[  ]"),
    ("Caminhos e nomes de arquivos atualizados no plano", "[  ]"),
    ("Página vazia removida do trabalho escrito", "[  ]"),
    ("Deck HTML testado no computador/projetor e PDF de backup salvo", "[  ]"),
    ("Cada integrante treinou sua parte e sabe explicar uma limitação do projeto", "[  ]"),
], [5.2, 1.2])

h1(doc, "9. Arquivos importantes do repositório")
table(doc, ["Arquivo", "Finalidade"], [
    ("docs/dw_canada.docx", "Trabalho escrito principal da Entrega 1."),
    ("presentation/index.html", "Apresentação interativa em formato de slides."),
    ("docs/dicionario_oltp.md", "Dicionário das fontes operacionais de origem."),
    ("docs/eda_oltp.md", "Roteiro e resultados da análise exploratória."),
    ("docs/modelo_dimensional.md", "Modelo Star Schema, grão e mapeamento origem-destino."),
    ("docs/dicionario_dw.md", "Dicionário das dimensões, fatos e staging."),
    ("sql/script_oracle.sql", "DDL Oracle a ser corrigido e anexado/executado."),
], [2.2, 4.2])

add_text(doc, "Referência didática: o exemplo de trabalho sobre NBA foi usado apenas como orientação de estrutura. Conforme aviso do docente, ele pode conter erros ou lacunas e não substitui os requisitos oficiais desta entrega.", size=9.5, italic=True, before=15, after=0)

doc.save(OUT)
print(OUT)
